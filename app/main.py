"""Zibris Cloud Director Inventory Web UI."""

from __future__ import annotations

import asyncio
import csv
import io
import os
import secrets
from datetime import datetime
from typing import Any, Optional
from zoneinfo import ZoneInfo

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

from app.vcd_client import VcdClient, VcdError, flatten_rows, normalize_inventory

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_BASE_URL = os.environ.get("VCD_BASE_URL", "https://portal.zibris-cloud.ch")
DEFAULT_API_VERSION = os.environ.get("VCD_API_VERSION", "39.1")
TZ = ZoneInfo(os.environ.get("TZ", "Europe/Zurich"))
SESSION_SECRET = os.environ.get("SESSION_SECRET") or secrets.token_hex(32)
SESSION_HTTPS_ONLY = os.environ.get("SESSION_HTTPS_ONLY", "false").lower() in {
    "1",
    "true",
    "yes",
    "on",
}

_INVENTORY_CACHE: dict[str, dict[str, Any]] = {}

app = FastAPI(title="Zibris VCD Inventory")
app.add_middleware(
    SessionMiddleware,
    secret_key=SESSION_SECRET,
    session_cookie="zibris_vcd_sess_v2",
    same_site="lax",
    https_only=SESSION_HTTPS_ONLY,
    max_age=60 * 60 * 8,
)
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "templates"))


class LoginBody(BaseModel):
    username: str = Field(min_length=1)
    org: str = Field(min_length=1)
    password: str = Field(min_length=1)
    base_url: str = DEFAULT_BASE_URL
    include_hdd_assigned: bool = False


def _session_client(request: Request) -> VcdClient:
    sess = request.session
    token = sess.get("token")
    if not token:
        raise HTTPException(status_code=401, detail="Nicht angemeldet.")
    client = VcdClient(
        base_url=sess.get("base_url") or DEFAULT_BASE_URL,
        username=sess.get("username") or "",
        org=sess.get("org") or "",
        password="",
        api_version=sess.get("api_version") or DEFAULT_API_VERSION,
        timeout=30.0,
    )
    client.token = token
    return client


def _cache_key(request: Request) -> str:
    key = request.session.get("cache_key")
    if not key:
        key = secrets.token_hex(16)
        request.session["cache_key"] = key
    return key


def _get_cached_inventory(request: Request) -> Optional[list[dict[str, Any]]]:
    key = request.session.get("cache_key")
    if not key:
        return None
    entry = _INVENTORY_CACHE.get(key)
    return entry["inventory"] if entry else None


def _now_label() -> str:
    return datetime.now(TZ).strftime("%d.%m.%Y %H:%M")


def _set_cached_inventory(request: Request, inventory: list[dict[str, Any]]) -> None:
    key = _cache_key(request)
    stamp = _now_label()
    _INVENTORY_CACHE[key] = {"inventory": inventory, "at": stamp}
    request.session["inventory_at"] = stamp
    request.session["count_vapps"] = len(inventory)
    request.session["count_vms"] = sum(len(v.get("vms") or []) for v in inventory)


def _clear_cache(request: Request) -> None:
    key = request.session.get("cache_key")
    if key:
        _INVENTORY_CACHE.pop(key, None)


def _fetch_inventory(
    client: VcdClient,
    include_hdd_assigned: bool = False,
) -> list[dict[str, Any]]:
    vapps = client.query_all("vApp")
    vms = client.query_all("vm", filter_expr="isVAppTemplate==false")
    inventory = normalize_inventory(vapps, vms)
    if include_hdd_assigned:
        return client.enrich_assigned_disks(inventory)
    return inventory


@app.get("/", response_class=HTMLResponse)
async def index(request: Request) -> HTMLResponse:
    # Always show the login form – do not auto-restore the previous session
    return templates.TemplateResponse(
        request,
        "index.html",
        {
            "default_base_url": DEFAULT_BASE_URL,
            "logged_in": False,
            "username": request.session.get("username") or "",
            "org": request.session.get("org") or "",
        },
    )


@app.get("/api/session")
async def session_info(request: Request) -> dict[str, Any]:
    cached = _get_cached_inventory(request)
    return {
        "logged_in": bool(request.session.get("token")),
        "has_cache": cached is not None,
        "username": request.session.get("username"),
        "org": request.session.get("org"),
        "base_url": request.session.get("base_url"),
        "inventory_at": request.session.get("inventory_at"),
        "count_vapps": request.session.get("count_vapps") or 0,
        "count_vms": request.session.get("count_vms") or 0,
    }


@app.post("/api/login")
async def login(body: LoginBody, request: Request) -> dict[str, Any]:
    client = VcdClient(
        base_url=(body.base_url or DEFAULT_BASE_URL).strip(),
        username=body.username.strip(),
        org=body.org.strip(),
        password=body.password,
        api_version=DEFAULT_API_VERSION,
        timeout=30.0,
    )

    def _do_login() -> tuple[dict[str, Any], list[dict[str, Any]]]:
        session = client.login()
        inventory = _fetch_inventory(
            client,
            include_hdd_assigned=bool(body.include_hdd_assigned),
        )
        return session, inventory

    try:
        session, inventory = await asyncio.to_thread(_do_login)
    except VcdError as exc:
        raise HTTPException(status_code=exc.status_code or 400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=f"Unerwarteter Fehler: {exc}") from exc

    _clear_cache(request)
    request.session.clear()
    request.session["token"] = client.token
    request.session["username"] = body.username.strip()
    request.session["org"] = body.org.strip()
    request.session["base_url"] = (body.base_url or DEFAULT_BASE_URL).strip()
    request.session["api_version"] = DEFAULT_API_VERSION
    _set_cached_inventory(request, inventory)

    return {
        "ok": True,
        "user": session.get("user", {}).get("name"),
        "org": session.get("org", {}).get("name"),
        "roles": session.get("roles") or [],
        "vapps": len(inventory),
        "vms": sum(len(v["vms"]) for v in inventory),
        "inventory_at": request.session.get("inventory_at"),
        "inventory": inventory,
    }


@app.post("/api/logout")
async def logout(request: Request) -> dict[str, bool]:
    _clear_cache(request)
    request.session.clear()
    return {"ok": True}


def _inventory_has_disk_api(inventory: list[dict[str, Any]] | None) -> bool:
    if not inventory:
        return False
    has_disks = False
    has_metrics_pass = False
    for vapp in inventory:
        if vapp.get("hdd_metrics_checked"):
            has_metrics_pass = True
        for vm in vapp.get("vms") or []:
            if vm.get("hdd_assigned_from_api") and (vm.get("hdd_disks") or []):
                has_disks = True
    return has_disks and has_metrics_pass


@app.get("/api/inventory")
async def get_inventory(
    request: Request,
    refresh: bool = False,
    include_hdd_assigned: bool = False,
) -> dict[str, Any]:
    client = _session_client(request)
    cached = _get_cached_inventory(request)

    cache_ok = cached is not None and (
        not include_hdd_assigned or _inventory_has_disk_api(cached)
    )
    if cache_ok and not refresh:
        return {
            "inventory": cached,
            "inventory_at": request.session.get("inventory_at"),
            "cached": True,
            "include_hdd_assigned": include_hdd_assigned,
        }

    if cached is None and not refresh and not include_hdd_assigned:
        raise HTTPException(
            status_code=401,
            detail="Session abgelaufen – bitte neu anmelden.",
        )

    try:
        data = await asyncio.to_thread(
            _fetch_inventory,
            client,
            bool(include_hdd_assigned),
        )
    except VcdError as exc:
        if exc.status_code in (401, 403):
            _clear_cache(request)
            request.session.clear()
        raise HTTPException(status_code=exc.status_code or 400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=500, detail=f"Abruf fehlgeschlagen: {exc}") from exc

    _set_cached_inventory(request, data)
    return {
        "inventory": data,
        "inventory_at": request.session.get("inventory_at"),
        "cached": False,
        "include_hdd_assigned": bool(include_hdd_assigned),
    }


@app.get("/api/debug/disk-metrics")
async def debug_disk_metrics(request: Request, vdc: str = "", limit: int = 10) -> dict[str, Any]:
    """Probe VCD metrics for a few VMs (HDD in use). Requires active login session."""
    client = _session_client(request)
    cached = _get_cached_inventory(request)
    if cached is None:
        raise HTTPException(status_code=400, detail="Kein Inventory geladen – bitte neu anmelden.")

    filtered = _filtered_inventory(cached, vdc=vdc)
    samples: list[dict[str, Any]] = []
    for vapp in filtered:
        for vm in vapp.get("vms") or []:
            if not vm.get("href"):
                continue
            samples.append(vm)
            if len(samples) >= max(1, min(limit, 50)):
                break
        if len(samples) >= max(1, min(limit, 50)):
            break

    def _probe() -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for vm in samples:
            metrics = client.vm_disk_usage_metrics(vm.get("href") or "")
            used_kib = metrics.get("used_kib")
            prov_kib = metrics.get("provisioned_kib")
            out.append(
                {
                    "vm": vm.get("name"),
                    "vapp": vm.get("vapp"),
                    "vdc": vm.get("vdc"),
                    "hdd_assigned_gib": vm.get("hdd_assigned_gib"),
                    "hdd_allok_gib": vm.get("hdd_gib"),
                    "metrics_ok": bool(metrics.get("ok")),
                    "status_code": metrics.get("status_code"),
                    "detail": metrics.get("detail") or "",
                    "metric_names": metrics.get("metric_names") or [],
                    "used_gib": round(float(used_kib) / 1024.0 / 1024.0, 2)
                    if used_kib is not None
                    else None,
                    "provisioned_gib": round(float(prov_kib) / 1024.0 / 1024.0, 2)
                    if prov_kib is not None
                    else None,
                }
            )
        return out

    results = await asyncio.to_thread(_probe)
    ok_count = sum(1 for r in results if r.get("metrics_ok") and r.get("used_gib") is not None)
    return {
        "vdc_filter": vdc or None,
        "checked": len(results),
        "with_used_metric": ok_count,
        "available": ok_count > 0,
        "note": (
            "disk.used.latest = Storage-Belegung (Thin), nicht Gast-OS-Frei/Belegt."
            if ok_count
            else "Metriken nicht verfügbar (Hoster hat sie oft deaktiviert) → HDD in use bleibt leer."
        ),
        "vms": results,
    }


def _filtered_inventory(
    data: list[dict[str, Any]],
    vdc: str = "",
    status: str = "",
) -> list[dict[str, Any]]:
    """Same filters as the UI: VDC + status (vApp and VMs)."""
    vdc = (vdc or "").strip()
    status = (status or "").strip()
    result: list[dict[str, Any]] = []

    for item in data:
        if vdc and (item.get("vdc") or "") != vdc:
            continue

        vms = list(item.get("vms") or [])
        if status:
            # Keep only VMs with matching status
            vms = [vm for vm in vms if (vm.get("status") or "") == status]
            # Keep vApp if its status matches or matching VMs remain
            vapp_ok = (item.get("status") or "") == status
            if not vapp_ok and not vms:
                continue

        filtered = dict(item)
        filtered["vms"] = vms
        filtered["vm_count"] = len(vms)
        if vms:
            filtered["cpu"] = sum(int(vm.get("cpu") or 0) for vm in vms)
            filtered["ram_gib"] = round(sum(float(vm.get("ram_gib") or 0) for vm in vms), 1)
            filtered["hdd_gib"] = round(sum(float(vm.get("hdd_gib") or 0) for vm in vms), 1)
            filtered["hdd_assigned_gib"] = round(
                sum(float(vm.get("hdd_assigned_gib") or 0) for vm in vms), 1
            )
            used_vms = [vm for vm in vms if vm.get("hdd_used_from_api")]
            filtered["hdd_used_from_api"] = bool(used_vms)
            filtered["hdd_used_gib"] = (
                round(sum(float(vm.get("hdd_used_gib") or 0) for vm in used_vms), 1)
                if used_vms
                else 0.0
            )
        result.append(filtered)

    return result


def _safe_filename_part(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in value.strip())
    return cleaned.strip("-")[:48] or "all"


def _export_filename(prefix: str, vdc: str, status: str, ext: str) -> str:
    parts = [prefix]
    parts.append(_safe_filename_part(vdc) if vdc.strip() else "alle-vdcs")
    if status.strip():
        short = {
            "POWERED_ON": "on",
            "POWERED_OFF": "off",
            "SUSPENDED": "sus",
            "RESOLVED": "res",
        }.get(status.strip(), _safe_filename_part(status))
        parts.append(short)
    parts.append(datetime.now().strftime("%Y%m%d-%H%M%S"))
    return f"{'-'.join(parts)}.{ext}"


def _parse_price(value: float | str | None) -> float:
    try:
        n = float(str(value if value is not None else "0").replace(",", "."))
    except (TypeError, ValueError):
        return 0.0
    return n if n >= 0 else 0.0


def _inventory_totals(filtered: list[dict[str, Any]]) -> dict[str, float]:
    cpu = 0.0
    ram = 0.0
    hdd = 0.0
    hdd_assigned = 0.0
    hdd_used = 0.0
    hdd_used_any = False
    vms = 0
    # Assumption: 1 external IP per VDC
    ips = len({(item.get("vdc") or "").strip() for item in filtered if (item.get("vdc") or "").strip()})
    for item in filtered:
        vm_list = item.get("vms") or []
        if vm_list:
            vms += len(vm_list)
            cpu += sum(float(vm.get("cpu") or 0) for vm in vm_list)
            ram += sum(float(vm.get("ram_gib") or 0) for vm in vm_list)
            hdd += sum(float(vm.get("hdd_gib") or 0) for vm in vm_list)
            hdd_assigned += sum(float(vm.get("hdd_assigned_gib") or 0) for vm in vm_list)
            for vm in vm_list:
                if vm.get("hdd_used_from_api"):
                    hdd_used_any = True
                    hdd_used += float(vm.get("hdd_used_gib") or 0)
        else:
            cpu += float(item.get("cpu") or 0)
            ram += float(item.get("ram_gib") or 0)
            hdd += float(item.get("hdd_gib") or 0)
            hdd_assigned += float(item.get("hdd_assigned_gib") or 0)
            if item.get("hdd_used_from_api"):
                hdd_used_any = True
                hdd_used += float(item.get("hdd_used_gib") or 0)
    # HDD cost base = assigned, fallback to allocation
    hdd_billable = hdd_assigned if hdd_assigned > 0 else hdd
    return {
        "vapps": float(len(filtered)),
        "vms": float(vms),
        "cpu": cpu,
        "ram_gib": round(ram, 1),
        "hdd_gib": round(hdd, 1),
        "hdd_assigned_gib": round(hdd_assigned, 1),
        "hdd_used_gib": round(hdd_used, 1) if hdd_used_any else 0.0,
        "hdd_used_available": 1.0 if hdd_used_any else 0.0,
        "hdd_billable_gib": round(hdd_billable, 1),
        "ips": float(ips),
    }


def _cost_summary(
    filtered: list[dict[str, Any]],
    price_cpu: float,
    price_ram: float,
    price_hdd: float,
    price_ip: float = 0,
    price_edge: float = 0,
    include_edge: bool = False,
) -> dict[str, Any]:
    totals = _inventory_totals(filtered)
    # Optional Edge GW: 1 per VDC when enabled
    edges = totals["ips"] if include_edge else 0.0
    cost_cpu = totals["cpu"] * price_cpu
    cost_ram = totals["ram_gib"] * price_ram
    cost_hdd = totals["hdd_billable_gib"] * price_hdd
    cost_ip = totals["ips"] * price_ip
    cost_edge = edges * price_edge
    return {
        **totals,
        "edges": float(edges),
        "include_edge": include_edge,
        "price_cpu": price_cpu,
        "price_ram": price_ram,
        "price_hdd": price_hdd,
        "price_ip": price_ip,
        "price_edge": price_edge,
        "cost_cpu": round(cost_cpu, 2),
        "cost_ram": round(cost_ram, 2),
        "cost_hdd": round(cost_hdd, 2),
        "cost_ip": round(cost_ip, 2),
        "cost_edge": round(cost_edge, 2),
        "cost_total": round(cost_cpu + cost_ram + cost_hdd + cost_ip + cost_edge, 2),
    }


def _parse_bool(value: str | bool | int | None) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "on"}


@app.get("/api/export.csv")
async def export_csv(
    request: Request,
    vdc: str = "",
    status: str = "",
    price_cpu: float = 0,
    price_ram: float = 0,
    price_hdd: float = 0,
    price_ip: float = 0,
    price_edge: float = 0,
    include_edge: str = "0",
) -> StreamingResponse:
    _session_client(request)
    data = _get_cached_inventory(request)
    if data is None:
        raise HTTPException(status_code=400, detail="Kein Inventory geladen – bitte neu anmelden.")
    filtered = _filtered_inventory(data, vdc=vdc, status=status)
    rows = flatten_rows(filtered)
    summary = _cost_summary(
        filtered,
        _parse_price(price_cpu),
        _parse_price(price_ram),
        _parse_price(price_hdd),
        _parse_price(price_ip),
        _parse_price(price_edge),
        _parse_bool(include_edge),
    )
    buf = io.StringIO()
    fieldnames = [
        "vapp",
        "vapp_status",
        "vdc",
        "vm",
        "vm_status",
        "os",
        "cpu",
        "ram_gib",
        "hdd_assigned_gib",
        "hdd_disks",
        "hdd_used_gib",
        "network",
    ]
    writer = csv.DictWriter(buf, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    buf.write("\n")
    buf.write("VM Berechnung\n")
    buf.write(f"Preis CPU (CHF/vCPU);{summary['price_cpu']}\n")
    buf.write(f"Preis RAM (CHF/GiB);{summary['price_ram']}\n")
    buf.write(f"Preis HDD (CHF/GiB zugewiesen);{summary['price_hdd']}\n")
    buf.write(f"Preis IP (CHF/VDC);{summary['price_ip']}\n")
    buf.write(f"Preis Edge GW (CHF/VDC);{summary['price_edge']}\n")
    buf.write(f"Edge GW einbezogen;{'ja' if summary['include_edge'] else 'nein'}\n")
    buf.write(f"Menge CPU (vCPU);{summary['cpu']}\n")
    buf.write(f"Menge RAM (GiB);{summary['ram_gib']}\n")
    buf.write(f"Menge HDD zugewiesen (GiB);{summary['hdd_assigned_gib']}\n")
    used_label = summary["hdd_used_gib"] if summary.get("hdd_used_available") else "n/a"
    buf.write(f"Menge HDD in use (GiB);{used_label}\n")
    buf.write(f"Menge IP (1 pro VDC);{int(summary['ips'])}\n")
    buf.write(f"Menge Edge GW (1 pro VDC);{int(summary['edges'])}\n")
    buf.write(f"Kosten CPU (CHF);{summary['cost_cpu']}\n")
    buf.write(f"Kosten RAM (CHF);{summary['cost_ram']}\n")
    buf.write(f"Kosten HDD (CHF);{summary['cost_hdd']}\n")
    buf.write(f"Kosten IP (CHF);{summary['cost_ip']}\n")
    buf.write(f"Kosten Edge GW (CHF);{summary['cost_edge']}\n")
    buf.write(f"Gesamtkosten (CHF);{summary['cost_total']}\n")
    content = buf.getvalue().encode("utf-8-sig")
    filename = _export_filename("zibris-vcd", vdc, status, "csv")
    return StreamingResponse(
        io.BytesIO(content),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/export.xlsx")
async def export_xlsx(
    request: Request,
    vdc: str = "",
    status: str = "",
    price_cpu: float = 0,
    price_ram: float = 0,
    price_hdd: float = 0,
    price_ip: float = 0,
    price_edge: float = 0,
    include_edge: str = "0",
) -> StreamingResponse:
    _session_client(request)
    data = _get_cached_inventory(request)
    if data is None:
        raise HTTPException(status_code=400, detail="Kein Inventory geladen – bitte neu anmelden.")
    filtered = _filtered_inventory(data, vdc=vdc, status=status)
    rows = flatten_rows(filtered)
    summary = _cost_summary(
        filtered,
        _parse_price(price_cpu),
        _parse_price(price_ram),
        _parse_price(price_hdd),
        _parse_price(price_ip),
        _parse_price(price_edge),
        _parse_bool(include_edge),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    headers = [
        "vApp",
        "vApp Status",
        "VDC",
        "VM",
        "VM Status",
        "OS",
        "CPU",
        "RAM (GiB)",
        "HDD zugewiesen (GiB)",
        "Festplatten (einzeln)",
        "HDD in use (GiB)",
        "Netzwerk",
    ]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="0F3D3E")
    header_font = Font(color="F4F7F5", bold=True)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(
            [
                row["vapp"],
                row["vapp_status"],
                row["vdc"],
                row["vm"],
                row["vm_status"],
                row["os"],
                row["cpu"],
                row["ram_gib"],
                row["hdd_assigned_gib"],
                row.get("hdd_disks") or "",
                row.get("hdd_used_gib", ""),
                row["network"],
            ]
        )

    data_end = ws.max_row
    widths = [28, 14, 24, 28, 14, 40, 8, 12, 16, 44, 16, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    if data_end >= 1:
        ws.auto_filter.ref = f"A1:L{data_end}"
    ws.freeze_panes = "A2"

    # Cost block at the bottom
    start = data_end + 2
    section_fill = PatternFill("solid", fgColor="0F3D3E")
    section_font = Font(color="F4F7F5", bold=True)
    label_font = Font(bold=True)
    total_fill = PatternFill("solid", fgColor="C45C26")
    total_font = Font(color="F4F7F5", bold=True)

    ws.cell(start, 1, "VM Berechnung").font = section_font
    ws.cell(start, 1).fill = section_fill
    ws.cell(start, 2).fill = section_fill
    ws.cell(start, 3).fill = section_fill

    scope_bits = []
    if vdc.strip():
        scope_bits.append(f"VDC: {vdc.strip()}")
    else:
        scope_bits.append("alle VDCs")
    if status.strip():
        scope_bits.append(status.strip())
    scope_bits.append(f"{int(summary['vapps'])} vApps · {int(summary['vms'])} VMs")
    if summary["include_edge"]:
        scope_bits.append("inkl. Edge GW")
    ws.cell(start + 1, 1, "Filter")
    ws.cell(start + 1, 2, " · ".join(scope_bits))

    price_rows = [
        (start + 3, "Preis CPU (CHF / vCPU)", summary["price_cpu"]),
        (start + 4, "Preis RAM (CHF / GiB)", summary["price_ram"]),
        (start + 5, "Preis HDD (CHF / GiB zugewiesen)", summary["price_hdd"]),
        (start + 6, "Preis IP (CHF / VDC)", summary["price_ip"]),
        (start + 7, "Preis Edge GW (CHF / VDC)", summary["price_edge"]),
    ]
    qty_rows = [
        (start + 9, "Menge CPU (vCPU)", summary["cpu"]),
        (start + 10, "Menge RAM (GiB)", summary["ram_gib"]),
        (start + 11, "Menge HDD zugewiesen (GiB)", summary["hdd_assigned_gib"]),
        (
            start + 12,
            "Menge HDD in use (GiB)",
            summary["hdd_used_gib"] if summary.get("hdd_used_available") else "n/a",
        ),
        (start + 13, "Menge IP (1 pro VDC)", int(summary["ips"])),
        (
            start + 14,
            "Menge Edge GW (1 pro VDC)",
            int(summary["edges"]),
        ),
    ]
    cost_rows = [
        (start + 16, "Kosten CPU (CHF)", summary["cost_cpu"]),
        (start + 17, "Kosten RAM (CHF)", summary["cost_ram"]),
        (start + 18, "Kosten HDD (CHF)", summary["cost_hdd"]),
        (start + 19, "Kosten IP (CHF)", summary["cost_ip"]),
        (start + 20, "Kosten Edge GW (CHF)", summary["cost_edge"]),
        (start + 21, "Gesamtkosten (CHF)", summary["cost_total"]),
    ]

    for r, label, value in price_rows + qty_rows + cost_rows:
        ws.cell(r, 1, label).font = label_font
        cell = ws.cell(r, 2, value)
        if "CHF" in label and "Preis" not in label:
            cell.number_format = "#,##0.00"
        elif "Preis" in label:
            cell.number_format = "#,##0.00"

    total_cell_label = ws.cell(start + 21, 1)
    total_cell_value = ws.cell(start + 21, 2)
    total_cell_label.fill = total_fill
    total_cell_label.font = total_font
    total_cell_value.fill = total_fill
    total_cell_value.font = total_font
    total_cell_value.number_format = "#,##0.00"

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = _export_filename("zibris-vcd", vdc, status, "xlsx")
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/healthz")
async def healthz() -> dict[str, str]:
    return {"status": "ok"}
